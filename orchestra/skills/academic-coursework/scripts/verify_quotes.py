#!/usr/bin/env python3
"""
Verify Chinese curly quote direction in a .docx file.
Uses pair-matching algorithm per paragraph:
  - First curly quote → should be LEFT (U+201C, " )
  - Second curly quote → should be RIGHT (U+201D, " )
  - Alternates thereafter

Reports misdirection errors and overall balance.
Exit code 0 = OK, 1 = errors found.

--fix mode now uses run-level character replacement, preserving all
formatting (bold, italic, font size, etc.) by only rewriting the
specific character in the specific run where it lives.
"""
import sys
from docx import Document

LEFT_Q = '“'   # "
RIGHT_Q = '”'  # "

def is_cjk(cp):
    return (0x4E00 <= cp <= 0x9FFF or 0x3400 <= cp <= 0x4DBF or
            0xF900 <= cp <= 0xFAFF or 0x3000 <= cp <= 0x303F or
            0xFF00 <= cp <= 0xFFEF or 0x2E80 <= cp <= 0x2FDF or
            0xFE30 <= cp <= 0xFE4F)

def verify_paragraph_text(text, label=""):
    """Verify quote direction in text. Returns list of error descriptions."""
    errors = []
    inside = False
    for i, ch in enumerate(text):
        if ch == LEFT_Q or ch == RIGHT_Q:
            if not inside:
                expected = LEFT_Q
                inside = True
            else:
                expected = RIGHT_Q
                inside = False

            if ch != expected:
                ctx = text[max(0,i-4):i+8].replace('\n', ' ')
                direction = "LEFT/opening" if expected == LEFT_Q else "RIGHT/closing"
                found = "LEFT" if ch == LEFT_Q else "RIGHT"
                errors.append(
                    f"{label} pos={i}: expected {direction}, got {found}  "
                    f"context: ...{ctx}..."
                )
    return errors


def fix_paragraph_runs(para):
    """Fix quote direction in a paragraph using run-level replacement.
    Preserves all run formatting (bold, italic, fonts, etc.).
    Returns number of characters fixed."""
    runs = para.runs
    if not runs:
        return 0

    # Build position map: char_offset → (run_idx, char_idx_in_run)
    pos_map = []  # [(run_idx, char_idx_in_run)]
    for ri, run in enumerate(runs):
        for ci, ch in enumerate(run.text):
            pos_map.append((ri, ci))

    if not pos_map:
        return 0

    full_text = para.text

    # Pair-matching scan: find all wrongly-directed quotes
    fixes = []  # [(offset_in_full_text, correct_char)]
    inside = False
    for i, ch in enumerate(full_text):
        if ch in (LEFT_Q, RIGHT_Q):
            correct = LEFT_Q if not inside else RIGHT_Q
            inside = not inside
            if ch != correct:
                fixes.append((i, correct))

    # Apply fixes back to runs
    fixed_count = 0
    for offset, correct_char in fixes:
        if offset < len(pos_map):
            ri, ci = pos_map[offset]
            run = runs[ri]
            text = run.text
            if ci < len(text) and text[ci] in (LEFT_Q, RIGHT_Q):
                run.text = text[:ci] + correct_char + text[ci+1:]
                fixed_count += 1

    return fixed_count


def fix_cell_paragraphs(cell):
    """Fix all paragraphs within a table cell."""
    total = 0
    for para in cell.paragraphs:
        total += fix_paragraph_runs(para)
    return total


def count_quotes(text, left_q, right_q):
    """Return (left_count, right_count) for a text."""
    return text.count(left_q), text.count(right_q)


def main():
    if len(sys.argv) < 2:
        print("Usage: python verify_quotes.py <file.docx> [--fix] [--sample N]")
        print("  --fix       Apply pair-matching fix (run-level, preserves formatting)")
        print("  --sample N  Only process first N paragraphs for verification")
        sys.exit(2)

    filepath = sys.argv[1]
    fix_mode = '--fix' in sys.argv
    sample_limit = None

    # Parse --sample N
    for i, arg in enumerate(sys.argv):
        if arg == '--sample' and i + 1 < len(sys.argv):
            sample_limit = int(sys.argv[i+1])
            break

    doc = Document(filepath)

    all_errors = []
    left_count = 0
    right_count = 0
    total_fixed = 0

    # Check and fix paragraphs
    para_count = len(doc.paragraphs)
    limit = min(sample_limit, para_count) if sample_limit else para_count

    if sample_limit:
        print(f"[SAMPLE MODE] Processing first {limit}/{para_count} paragraphs")

    for pi, para in enumerate(doc.paragraphs):
        text = para.text
        lc, rc = count_quotes(text, LEFT_Q, RIGHT_Q)
        left_count += lc
        right_count += rc

        if pi < limit:
            errors = verify_paragraph_text(text, f"P{pi}")
            all_errors.extend(errors)

            if fix_mode and text:
                fixed = fix_paragraph_runs(para)
                total_fixed += fixed

    # Check and fix table cells (tables are never sampled for now)
    for ti, table in enumerate(doc.tables):
        for ri, row in enumerate(table.rows):
            for ci, cell in enumerate(row.cells):
                text = cell.text
                lc, rc = count_quotes(text, LEFT_Q, RIGHT_Q)
                left_count += lc
                right_count += rc

                if not sample_limit:
                    errors = verify_paragraph_text(text, f"T{ti}R{ri}C{ci}")
                    all_errors.extend(errors)

                    if fix_mode and text:
                        fixed = fix_cell_paragraphs(cell)
                        total_fixed += fixed

    # Report
    print(f"Chinese curly quotes: LEFT={left_count}, RIGHT={right_count}")
    imbalance = abs(left_count - right_count)
    print(f"Imbalance: {imbalance}")

    if all_errors:
        print(f"\n{len(all_errors)} direction errors found:")
        for e in all_errors[:20]:
            print(f"  {e}")
        if len(all_errors) > 20:
            print(f"  ... and {len(all_errors) - 20} more")

    if fix_mode and (all_errors or imbalance > 2):
        doc.save(filepath)
        print(f"\n[FIXED] {total_fixed} characters corrected via run-level replacement")
        print(f"  Saved to: {filepath}")

    if all_errors:
        sys.exit(1)
    if imbalance > 2:
        print("\nWARNING: Imbalance > 2 may indicate unclosed quotes or multi-paragraph quotations")
        sys.exit(1)
    print("OK")
    sys.exit(0)


if __name__ == '__main__':
    main()
